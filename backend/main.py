
import os
import re
import traceback
import uuid
from pathlib import Path

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from sentence_transformers import SentenceTransformer
import chromadb
from google import genai


# --------------------------------------------------
# 1. SETTINGS
# --------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-flash-lite-latest",
)
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

app = FastAPI(title="KnowSphere AI")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# Keep using your existing database.
embedding_model = SentenceTransformer(
    "all-MiniLM-L6-v2"
)

chroma_client = chromadb.PersistentClient(
    path=str(BASE_DIR / "chroma_db")
)

collection = chroma_client.get_or_create_collection(
    name="knowsphere_documents"
)

gemini_client = (
    genai.Client(api_key=GEMINI_API_KEY)
    if GEMINI_API_KEY
    else None
)


# --------------------------------------------------
# 2. REQUEST MODEL
# --------------------------------------------------

class QuestionRequest(BaseModel):
    question: str


# --------------------------------------------------
# 3. TEXT PROCESSING
# --------------------------------------------------

def chunk_text(
    text: str,
    chunk_size: int = 500,
) -> list[str]:
    """Split text into word-based chunks."""
    words = text.split()

    if not words:
        return []

    return [
        " ".join(words[start:start + chunk_size])
        for start in range(0, len(words), chunk_size)
    ]


def extract_text(file_path: Path) -> str:
    """Extract text from TXT, PDF, or DOCX files."""
    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        return file_path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    if suffix == ".docx":
        document = Document(str(file_path))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)

    raise HTTPException(
        status_code=400,
        detail=(
            "Unsupported file type. "
            "Upload a PDF, DOCX, or TXT file."
        ),
    )


# --------------------------------------------------
# 4. IMPROVED DOCUMENT RETRIEVAL
# --------------------------------------------------


STOP_WORDS = {
    "what", "is", "are", "was", "were", "the",
    "a", "an", "of", "in", "on", "to", "for",
    "and", "or", "but", "me", "my", "mine",
    "our", "your", "their", "his", "her",
    "explain", "tell", "about", "please", "define",
    "does", "do", "did", "how", "why", "can",
    "could", "would", "should", "will", "you",
    "be", "been", "being", "it", "this", "that",
    "from", "with", "which", "who", "whom",
    "where", "when", "i", "we", "he", "she",
    "they", "have", "has", "had", "any",
    "information", "document", "documents",
}


def get_query_terms(question: str) -> set[str]:
    """Extract meaningful search terms."""
    words = re.findall(
        r"[a-zA-Z0-9]+",
        question.lower(),
    )

    return {
        word
        for word in words
        if word not in STOP_WORDS
    }



def retrieve_chunks(
    question: str,
    limit: int = 5,
) -> list[dict]:
    """
    Retrieve relevant document chunks using semantic similarity
    and keyword matching. Avoid returning unrelated documents.
    """
    total_chunks = collection.count()

    if total_chunks == 0:
        return []

    query_embedding = embedding_model.encode(
        question
    ).tolist()

    candidate_count = min(
        total_chunks,
        max(limit * 4, 10),
    )

    result = collection.query(
        query_embeddings=[query_embedding],
        n_results=candidate_count,
        include=[
            "documents",
            "metadatas",
            "distances",
        ],
    )

    ids = (result.get("ids") or [[]])[0]
    documents = (result.get("documents") or [[]])[0]
    metadatas = (result.get("metadatas") or [[]])[0]
    distances = (result.get("distances") or [[]])[0]

    query_terms = get_query_terms(question)
    chunks = []
    seen_text = set()

    for index, text in enumerate(documents):
        if not isinstance(text, str) or not text.strip():
            continue

        cleaned_text = text.strip()

        if cleaned_text in seen_text:
            continue
        seen_text.add(cleaned_text)

        metadata = (
            metadatas[index]
            if index < len(metadatas)
            else None
        ) or {}

        filename = metadata.get("filename", "Unknown file")
        distance = (
            distances[index]
            if index < len(distances)
            else None
        )

        filename_terms = set(
            re.findall(
                r"[a-zA-Z0-9]+",
                Path(filename).stem.lower(),
            )
        )
        text_terms = set(
            re.findall(
                r"[a-zA-Z0-9]+",
                cleaned_text.lower(),
            )
        )

        filename_matches = len(query_terms & filename_terms)
        text_matches = len(query_terms & text_terms)

        score = (
            -float(distance)
            if distance is not None
            else 0.0
        )
        score += 2.5 * filename_matches
        score += 0.35 * text_matches

        chunks.append(
            {
                "id": (
                    ids[index]
                    if index < len(ids)
                    else str(uuid.uuid4())
                ),
                "filename": filename,
                "chunk_index": metadata.get("chunk_index"),
                "text": cleaned_text,
                "distance": distance,
                "score": score,
                "filename_matches": filename_matches,
                "text_matches": text_matches,
            }
        )

    chunks.sort(
        key=lambda chunk: chunk["score"],
        reverse=True,
    )

    
    # Prefer chunks that match meaningful question terms.
    # If the question includes a number such as employee ID 202,
    # require that number to appear in the document chunk.
    numeric_query_terms = {
        term for term in query_terms if term.isdigit()
    }

    relevant_chunks = []

    for chunk in chunks:
        text_terms = set(
            re.findall(
                r"[a-zA-Z0-9]+",
                chunk["text"].lower(),
            )
        )

        # Avoid matching a different employee or record
        # merely because it contains generic words.
        if numeric_query_terms and not (
            numeric_query_terms & text_terms
        ):
            continue

        if (
            chunk["filename_matches"] > 0
            or chunk["text_matches"] > 0
        ):
            relevant_chunks.append(chunk)

    # If no keyword matches remain, use close semantic matches
    # only when the question does not contain a numeric ID.
    if not relevant_chunks and not numeric_query_terms:
        relevant_chunks = [
            chunk
            for chunk in chunks
            if (
                chunk["distance"] is not None
                and float(chunk["distance"]) <= 1.30
            )
        ]

    print("DEBUG: Query terms:", query_terms)
    print(
        "DEBUG: Matching results:",
        [
            {
                "filename": chunk["filename"],
                "distance": chunk["distance"],
                "filename_matches": chunk["filename_matches"],
                "text_matches": chunk["text_matches"],
            }
            for chunk in relevant_chunks[:limit]
        ],
    )

    return relevant_chunks[:limit]

# --------------------------------------------------
# 5. HOME ROUTE
# --------------------------------------------------

@app.get("/")
def root():
    return {
        "message": "KnowSphere AI backend is running."
    }


# --------------------------------------------------
# 6. UPLOAD DOCUMENT
# --------------------------------------------------

@app.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
):
    filename = Path(
        file.filename or "uploaded_file"
    ).name

    suffix = Path(filename).suffix.lower()

    if suffix not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Upload a PDF, DOCX, or TXT file."
            ),
        )

    saved_path = UPLOAD_DIR / (
        f"{uuid.uuid4().hex}_{filename}"
    )

    try:
        contents = await file.read()

        if not contents:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty.",
            )

        saved_path.write_bytes(contents)

        extracted = extract_text(
            saved_path
        ).strip()

        if not extracted:
            raise HTTPException(
                status_code=400,
                detail="No readable text was found.",
            )

        chunks = chunk_text(extracted)

        if not chunks:
            raise HTTPException(
                status_code=400,
                detail="No usable text chunks were created.",
            )

        embeddings = embedding_model.encode(
            chunks
        ).tolist()

        document_id = uuid.uuid4().hex

        ids = [
            f"{document_id}_{index}"
            for index in range(len(chunks))
        ]

        metadatas = [
            {
                "filename": filename,
                "chunk_index": index,
                "document_id": document_id,
            }
            for index in range(len(chunks))
        ]

        collection.add(
            ids=ids,
            documents=chunks,
            embeddings=embeddings,
            metadatas=metadatas,
        )

        return {
            "message": (
                "Document uploaded and processed successfully."
            ),
            "filename": filename,
            "chunks_created": len(chunks),
            "document_id": document_id,
        }

    except HTTPException:
        # Remove a partially saved upload if processing failed.
        if saved_path.exists():
            saved_path.unlink()
        raise

    except Exception as exc:
        print(
            "KnowSphere AI /upload error:",
            repr(exc),
        )
        traceback.print_exc()

        if saved_path.exists():
            saved_path.unlink()

        raise HTTPException(
            status_code=500,
            detail=(
                f"Document upload failed: "
                f"{type(exc).__name__}."
            ),
        ) from exc

    finally:
        await file.close()


# --------------------------------------------------
# 7. LIST UPLOADED DOCUMENTS
# --------------------------------------------------

@app.get("/documents")
def list_documents():
    documents = []

    for path in UPLOAD_DIR.iterdir():
        if (
            path.is_file()
            and path.suffix.lower()
            in {".pdf", ".docx", ".txt"}
        ):
            documents.append(
                {
                    "filename": path.name.split(
                        "_", 1
                    )[-1],
                    "stored_name": path.name,
                    "size_bytes": path.stat().st_size,
                    "uploaded_at": path.stat().st_mtime,
                }
            )

    documents.sort(
        key=lambda document: document["uploaded_at"],
        reverse=True,
    )

    return {
        "documents": documents,
        "count": len(documents),
    }


# --------------------------------------------------
# 8. OPEN AN UPLOADED DOCUMENT
# --------------------------------------------------

@app.get("/documents/open/{filename}")
def open_document(filename: str):
    safe_name = Path(filename).name
    file_path = UPLOAD_DIR / safe_name

    if (
        not file_path.is_file()
        or file_path.parent != UPLOAD_DIR
    ):
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    return FileResponse(
        path=str(file_path),
        filename=safe_name,
        content_disposition_type="inline",
    )


# --------------------------------------------------
# 9. SEARCH DOCUMENTS
# --------------------------------------------------

@app.post("/search")
def search_documents(
    request: QuestionRequest,
):
    question = request.question.strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Please enter a search question.",
        )

    try:
        chunks = retrieve_chunks(
            question,
            limit=5,
        )

        return {
            "question": question,
            "results": [
                {
                    "id": chunk["id"],
                    "filename": chunk["filename"],
                    "chunk": (
                        chunk["chunk_index"] + 1
                        if isinstance(
                            chunk["chunk_index"], int
                        )
                        else None
                    ),
                    "text": chunk["text"],
                    "distance": chunk["distance"],
                }
                for chunk in chunks
            ],
            "result_count": len(chunks),
        }

    except HTTPException:
        raise

    except Exception as exc:
        print(
            "KnowSphere AI /search error:",
            repr(exc),
        )
        traceback.print_exc()

        raise HTTPException(
            status_code=500,
            detail=(
                f"Document search failed: "
                f"{type(exc).__name__}."
            ),
        ) from exc


# --------------------------------------------------
# 10. ASK AI
# --------------------------------------------------


@app.post("/ask")
def ask_ai(request: QuestionRequest):
    question = request.question.strip()

    if not question:
        raise HTTPException(
            status_code=400,
            detail="Please enter a question.",
        )

    if gemini_client is None:
        raise HTTPException(
            status_code=503,
            detail="Gemini API key is missing from backend/.env.",
        )

    refusal = (
        "I could not find enough information to answer "
        "this in the uploaded document."
    )

    try:
        chunks = retrieve_chunks(question, limit=5)

        print("DEBUG /ask question:", question)
        print("DEBUG /ask retrieved chunks:", len(chunks))

        if not chunks:
            return {
                "question": question,
                "answer": refusal,
                "sources": [],
                "source_count": 0,
            }

        context_parts = []

        for index, chunk in enumerate(chunks, start=1):
            context_parts.append(
                f"[Source {index}]\n"
                f"Filename: {chunk['filename']}\n"
                f"Content:\n{chunk['text']}"
            )

        context = "\n\n---\n\n".join(context_parts)

        prompt = f"""You are KnowSphere AI, a document-based AI assistant.

STRICT RULES:
- Answer using only facts supported by the supplied document context.
- Do not use outside knowledge to answer personal or factual questions.
- Do not guess, assume, or combine unrelated documents to invent an answer.
- If the context does not clearly support an answer, reply exactly:
  "{refusal}"
- Be concise and clear.
- Do not include source numbers in the answer.

DOCUMENT CONTEXT:
{context}

USER QUESTION:
{question}

ANSWER:
"""

        response = gemini_client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
        )

        answer = (
            getattr(response, "text", None)
            if response is not None
            else None
        )

        if not isinstance(answer, str) or not answer.strip():
            raise HTTPException(
                status_code=502,
                detail=(
                    "Gemini returned no text. Check the model, "
                    "quota, and API response."
                ),
            )

        answer = answer.strip()

        # If the AI refuses, no document is presented as a source
        # for an answer that could not be supported.
        normalized_answer = answer.lower().strip().strip('"').strip("'")

        refusal_phrases = [
            "i could not find enough information",
            "i couldn't find enough information",
            "there is not enough information",
            "there is insufficient information",
            "the documents do not contain enough information",
        ]

        is_refusal = any(
            phrase in normalized_answer
            for phrase in refusal_phrases
        )

        if is_refusal:
            answer = refusal
            sources = []
        else:
            # Keep sources with direct word overlap with the question.
            # The complete retrieved context is still given to Gemini,
            # so semantic retrieval can help answer paraphrased questions.
            relevant_chunks = [
                chunk
                for chunk in chunks
                if (
                    chunk.get("filename_matches", 0) > 0
                    or chunk.get("text_matches", 0) > 0
                )
            ]

            # A paraphrased question may not share exact keywords
            # with the document. In that case, use retrieved chunks
            # as a fallback so the answer can still show its evidence.
            if not relevant_chunks:
                relevant_chunks = chunks

            sources = [
                {
                    "filename": chunk["filename"],
                    "chunk": (
                        chunk["chunk_index"] + 1
                        if isinstance(chunk.get("chunk_index"), int)
                        else None
                    ),
                    "text": chunk["text"],
                }
                for chunk in relevant_chunks
            ]

        return {
            "question": question,
            "answer": answer,
            "sources": sources,
            "source_count": len(sources),
        }

    except HTTPException:
        raise

    except Exception as exc:
        print("KnowSphere AI /ask error:", repr(exc))
        traceback.print_exc()

        raise HTTPException(
            status_code=502,
            detail=(
                f"AI question answering failed: "
                f"{type(exc).__name__}."
            ),
        ) from exc
